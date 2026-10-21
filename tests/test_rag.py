from types import SimpleNamespace

import pytest

from app.models.document import DOCUMENT_STATUS_QUEUED, DOCUMENT_STATUS_READY, Document
from app.services.rag_service import RAGService


def _create_user_and_login(client, suffix: str) -> str:
    email = f"rag-{suffix}@example.com"
    password = "password123"
    client.post(
        "/api/v1/users/",
        json={
            "username": f"raguser-{suffix}",
            "email": email,
            "password": password,
        },
    )
    response = client.post(
        "/api/v1/auth/login",
        data={"username": email, "password": password},
    )
    assert response.status_code == 200
    return response.json()["access_token"]


def test_upload_document_creates_async_job(client, monkeypatch, db_session):
    token = _create_user_and_login(client, "upload")

    class DummyTask:
        id = "task-rag-upload-1"

    monkeypatch.setattr(
        "app.api.routers.rag.process_document_task.delay",
        lambda document_id: DummyTask(),
    )

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("notes.txt", b"hello rag", "text/plain")},
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload["status"] == DOCUMENT_STATUS_QUEUED
    assert payload["processing_task_id"] == "task-rag-upload-1"
    assert payload["chunks_count"] == 0
    assert payload["file_type"] == "txt"

    document = db_session.get(Document, payload["id"])
    assert document is not None
    assert document.status == DOCUMENT_STATUS_QUEUED
    assert document.processing_task_id == "task-rag-upload-1"
    assert document.owner_id is not None
    assert document.file_type == "txt"


def test_upload_docx_document(client, monkeypatch):
    token = _create_user_and_login(client, "upload-docx")

    class DummyTask:
        id = "task-rag-docx"

    monkeypatch.setattr(
        "app.api.routers.rag.process_document_task.delay",
        lambda document_id: DummyTask(),
    )

    import io

    from docx import Document as DocxDocument

    buf = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph("雅思写作 task 2 结构")
    document.save(buf)

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "ielts.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload["file_type"] == "docx"
    assert payload["status"] == DOCUMENT_STATUS_QUEUED


def test_upload_legacy_doc_returns_actionable_hint(client):
    """老格式必须给出"另存为 docx"这种可操作提示，而不是笼统的不支持。"""
    token = _create_user_and_login(client, "upload-legacy-doc")

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("ielts.doc", b"legacy binary", "application/msword")},
    )

    assert response.status_code == 400
    assert "另存为 .docx" in str(response.json())


def test_upload_rejects_oversized_file(client, monkeypatch):
    import app.api.routers.rag as rag_router

    token = _create_user_and_login(client, "upload-big")
    monkeypatch.setattr(rag_router, "MAX_UPLOAD_BYTES", 16)

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("big.txt", b"x" * 64, "text/plain")},
    )

    assert response.status_code == 413
    assert "文件过大" in str(response.json())


def test_query_knowledge_base_uses_current_user_scope(client, monkeypatch, db_session):
    token = _create_user_and_login(client, "query")
    user_id = int(
        client.get("/api/v1/users/me", headers={"Authorization": f"Bearer {token}"}).json()["id"]
    )

    ready_doc = Document(
        owner_id=user_id,
        filename="ready.txt",
        content="rag content",
        status=DOCUMENT_STATUS_READY,
        chunks_count=1,
    )
    db_session.add(ready_doc)
    db_session.commit()

    async def fake_retrieve(_self, *, query=None, owner_id, top_k=None):
        assert owner_id == user_id
        assert query == "我的知识库里有什么？"
        assert top_k == 2
        return [SimpleNamespace(content="只属于当前用户的知识片段")]

    def fake_get_llm_client(user=None):
        class DummyClient:
            class chat:
                class completions:
                    @staticmethod
                    async def create(**_kwargs):
                        return SimpleNamespace(
                            choices=[
                                SimpleNamespace(
                                    message=SimpleNamespace(content="基于当前用户文档的回答")
                                )
                            ],
                            usage=SimpleNamespace(
                                prompt_tokens=5, completion_tokens=7, total_tokens=12
                            ),
                        )

        return SimpleNamespace(client=DummyClient(), model_name="test-model")

    monkeypatch.setattr(RAGService, "retrieve_relevant_chunks", fake_retrieve)
    monkeypatch.setattr("app.api.routers.rag.get_llm_client", fake_get_llm_client)

    response = client.post(
        "/api/v1/rag/query",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "我的知识库里有什么？", "top_k": 2},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["answer"] == "基于当前用户文档的回答"
    assert payload["source_chunks"] == ["只属于当前用户的知识片段"]


def test_rag_query_global_question_injects_summaries(client, monkeypatch, db_session):
    """回归：全局性问题必须带上文档摘要兜底。

    线上真实故障：问"串讲四篇论文"时 top-3 检索捞到的是参考文献列表片段，
    模型只能看到零碎引用。摘要由 Celery beat 每晚生成，这里验证它真的被检索链路读到。
    """
    token = _create_user_and_login(client, "global-summary")
    user_id = int(
        client.get("/api/v1/users/me", headers={"Authorization": f"Bearer {token}"}).json()["id"]
    )

    db_session.add(
        Document(
            owner_id=user_id,
            filename="Transformer.pdf",
            file_type="pdf",
            content="正文",
            status=DOCUMENT_STATUS_READY,
            chunks_count=1,
            summary="提出自注意力机制，可并行训练",
        )
    )
    db_session.commit()

    async def fake_retrieve(_self, *, query, owner_id, top_k):
        # 模拟线上真实故障：全局问题一个原文片段都没检索到
        return []

    captured = {}

    def fake_get_llm_client(user=None):
        class DummyClient:
            class chat:
                class completions:
                    @staticmethod
                    async def create(**kwargs):
                        captured["messages"] = kwargs["messages"]
                        return SimpleNamespace(
                            choices=[SimpleNamespace(message=SimpleNamespace(content="串讲如下"))],
                            usage=SimpleNamespace(
                                prompt_tokens=5, completion_tokens=7, total_tokens=12
                            ),
                        )

        return SimpleNamespace(client=DummyClient(), model_name="test-model")

    monkeypatch.setattr(RAGService, "retrieve_relevant_chunks", fake_retrieve)
    monkeypatch.setattr("app.api.routers.rag.get_llm_client", fake_get_llm_client)

    response = client.post(
        "/api/v1/rag/query",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "帮我串讲这些论文", "top_k": 3},
    )

    assert response.status_code == 200
    system_prompt = captured["messages"][0]["content"]
    assert "文档摘要" in system_prompt
    assert "提出自注意力机制，可并行训练" in system_prompt
    assert any("（文档摘要）" in label for label in response.json()["source_chunks"])


def test_worker_requeues_owned_document_and_returns_task_status(client, monkeypatch, db_session):
    token = _create_user_and_login(client, "worker")
    user_id = int(
        client.get("/api/v1/users/me", headers={"Authorization": f"Bearer {token}"}).json()["id"]
    )

    document = Document(
        owner_id=user_id,
        filename="retry.txt",
        content="retry me",
        status=DOCUMENT_STATUS_READY,
        chunks_count=2,
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    class DummyTask:
        id = "task-rag-worker-1"

    class DummyAsyncResult:
        status = "PROGRESS"
        info = {"step": "processing_document", "current": 1, "total": 1}

    monkeypatch.setattr(
        "app.api.routers.worker.process_document_task.delay",
        lambda document_id: DummyTask(),
    )
    monkeypatch.setattr(
        "app.api.routers.worker.celery_app.AsyncResult",
        lambda task_id: DummyAsyncResult(),
    )

    process_response = client.post(
        "/api/v1/worker/process",
        headers={"Authorization": f"Bearer {token}"},
        json={"document_id": document.id},
    )
    assert process_response.status_code == 200
    process_payload = process_response.json()
    assert process_payload["task_id"] == "task-rag-worker-1"
    assert process_payload["status"] == DOCUMENT_STATUS_QUEUED

    db_session.expire_all()
    updated_document = db_session.get(Document, document.id)
    assert updated_document.processing_task_id == "task-rag-worker-1"
    assert updated_document.status == DOCUMENT_STATUS_QUEUED

    status_response = client.get(
        "/api/v1/worker/status/task-rag-worker-1",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert status_response.status_code == 200
    assert status_response.json()["result"] == {
        "step": "processing_document",
        "current": 1,
        "total": 1,
    }


def test_upload_md_document(client, monkeypatch, db_session):
    token = _create_user_and_login(client, "md-upload")

    class DummyTask:
        id = "task-rag-md-1"

    monkeypatch.setattr(
        "app.api.routers.rag.process_document_task.delay",
        lambda document_id: DummyTask(),
    )

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("readme.md", b"# Title\n\nSome markdown content.", "text/markdown")},
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload["file_type"] == "md"
    assert payload["status"] == DOCUMENT_STATUS_QUEUED

    document = db_session.get(Document, payload["id"])
    assert document.file_type == "md"
    assert "Title" in document.content


def test_upload_pdf_document(client, monkeypatch, db_session):
    pytest.importorskip("fpdf", reason="fpdf2 为测试依赖（requirements-dev.txt），未安装则跳过")
    token = _create_user_and_login(client, "pdf-upload")

    import io

    from fpdf import FPDF

    buf = io.BytesIO()
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(200, 10, text="Project Orion RAG Test", new_x="LMARGIN", new_y="NEXT")
    pdf.output(buf)
    pdf_bytes = buf.getvalue()

    class DummyTask:
        id = "task-rag-pdf-1"

    monkeypatch.setattr(
        "app.api.routers.rag.process_document_task.delay",
        lambda document_id: DummyTask(),
    )

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("report.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload["file_type"] == "pdf"
    assert payload["status"] == DOCUMENT_STATUS_QUEUED

    document = db_session.get(Document, payload["id"])
    assert document.file_type == "pdf"
    assert "Project Orion" in document.content


def test_upload_unsupported_format(client):
    token = _create_user_and_login(client, "bad-fmt")

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("image.png", b"fake png data", "image/png")},
    )

    assert response.status_code == 400
    assert "Unsupported" in response.json()["msg"]


def test_upload_empty_pdf(client):
    pytest.importorskip("fpdf", reason="fpdf2 为测试依赖（requirements-dev.txt），未安装则跳过")
    token = _create_user_and_login(client, "empty-pdf")

    import io

    from fpdf import FPDF

    buf = io.BytesIO()
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.output(buf)

    response = client.post(
        "/api/v1/rag/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("empty.pdf", buf.getvalue(), "application/pdf")},
    )

    assert response.status_code == 400
    assert "no extractable text" in response.json()["msg"].lower()


def test_rag_query_stream_returns_sse(client, monkeypatch):
    token = _create_user_and_login(client, "stream")

    class FakeChunk:
        content = "streamed reply chunk"

    async def fake_retrieve(_self, *, query=None, owner_id=None, top_k=None):
        return [FakeChunk()]

    async def fake_stream_create(**kwargs):
        async def gen():
            chunk = type(
                "C",
                (),
                {
                    "choices": [
                        type(
                            "Ch",
                            (),
                            {"delta": type("D", (), {"content": "Hello from RAG stream"})()},
                        )()
                    ]
                },
            )()
            yield chunk

        return gen()

    monkeypatch.setattr(RAGService, "retrieve_relevant_chunks", fake_retrieve)
    monkeypatch.setattr(
        "app.api.routers.rag.get_llm_client",
        lambda user=None: SimpleNamespace(
            client=type(
                "C",
                (),
                {
                    "chat": type(
                        "Ch",
                        (),
                        {
                            "completions": type(
                                "Co", (), {"create": staticmethod(fake_stream_create)}
                            )()
                        },
                    )()
                },
            )(),
            model_name="test-model",
        ),
    )

    with client.stream(
        "POST",
        "/api/v1/rag/query/stream",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "test", "top_k": 3},
    ) as resp:
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "text/event-stream; charset=utf-8"


def test_rag_query_stream_no_chunks(client, monkeypatch, db_session):
    token = _create_user_and_login(client, "stream-empty")

    async def fake_retrieve_empty(*a, **kw):
        return []

    monkeypatch.setattr(RAGService, "retrieve_relevant_chunks", fake_retrieve_empty)

    with client.stream(
        "POST",
        "/api/v1/rag/query/stream",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "test", "top_k": 3},
    ) as resp:
        assert resp.status_code == 200
        content = resp.read().decode("utf-8")
        assert "暂无已处理" in content or "data:" in content


def test_delete_document_endpoint(client, db_session):
    """删除接口：返回体要能拿到文件名（删除后实例已失效，属性访问极易踩坑）。"""
    token = _create_user_and_login(client, "delete-doc")
    user_id = int(
        client.get("/api/v1/users/me", headers={"Authorization": f"Bearer {token}"}).json()["id"]
    )

    document = Document(
        owner_id=user_id,
        filename="to-delete.md",
        file_type="md",
        content="待删除内容",
        status=DOCUMENT_STATUS_READY,
        chunks_count=0,
    )
    db_session.add(document)
    db_session.commit()
    document_id = document.id

    response = client.delete(
        f"/api/v1/rag/documents/{document_id}",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert "to-delete.md" in payload["message"]
    assert payload["document_id"] == document_id
    assert db_session.query(Document).filter_by(id=document_id).count() == 0


def test_rag_search_returns_chunks(client, db_session, monkeypatch):
    """纯检索接口：返回相关片段且不调用 LLM"""
    # 注册用户并拿到 id
    client.post(
        "/api/v1/users/",
        json={"username": "searcher", "email": "searcher@test.dev", "password": "testpassword"},
    )
    login = client.post(
        "/api/v1/auth/login",
        data={"username": "searcher@test.dev", "password": "testpassword"},
    )
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    # 通过 API 拿当前用户 id
    me = client.get("/api/v1/users/me", headers=headers).json()
    user_id = me["id"]

    from app.models.document import DOCUMENT_STATUS_READY, Document, DocumentChunk

    doc = Document(
        owner_id=user_id,
        filename="searchme.md",
        file_type="md",
        content="Faiss 是向量检索库" * 20,
        status=DOCUMENT_STATUS_READY,
    )
    db_session.add(doc)
    db_session.commit()

    async def fake_embeddings(self, texts):
        # 简单确定性向量：非空文本返回固定向量
        return [[1.0] * 1024 for _ in texts]

    monkeypatch.setattr("app.services.rag_service.RAGService.get_embeddings", fake_embeddings)
    db_session.add(
        DocumentChunk(
            document_id=doc.id,
            chunk_index=0,
            content="Faiss 是 Facebook 开源的向量相似度检索库",
            embedding=[1.0] * 1024,
        )
    )
    db_session.commit()

    resp = client.get("/api/v1/rag/search", params={"query": "向量检索库"}, headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["query"] == "向量检索库"
    assert len(data["results"]) >= 1
    assert "Faiss" in data["results"][0]["content"]


def test_rag_query_budget_blocked(client, db_session):
    """/rag/query 必须受成本熔断约束（此前绕过治理的洞）"""
    from app.core.config import settings
    from app.models.llm_call_log import LLMCallLog

    old_budget = settings.LLM_BUDGET_CNY
    settings.LLM_BUDGET_CNY = 0.0001
    client.post(
        "/api/v1/users/",
        json={"username": "ragbudget", "email": "ragbudget@test.dev", "password": "testpassword"},
    )
    login = client.post(
        "/api/v1/auth/login",
        data={"username": "ragbudget@test.dev", "password": "testpassword"},
    )
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    user_id = client.get("/api/v1/users/me", headers=headers).json()["id"]

    db_session.add(
        LLMCallLog(
            user_id=user_id,
            endpoint="/api/v1/chat",
            model_name="deepseek-v4.1-flash",
            provider="custom",
            prompt="hi",
            response="x",
            prompt_tokens=1_000_000,
            completion_tokens=0,
            total_tokens=1_000_000,
            latency_ms=1.0,
            estimated_cost_usd=0.0,
            status="success",
        )
    )
    db_session.commit()

    resp = client.post("/api/v1/rag/query", json={"query": "测试"}, headers=headers)
    assert resp.status_code == 402
    assert "成本熔断" in str(resp.json())
    settings.LLM_BUDGET_CNY = old_budget


def test_rag_query_stream_budget_blocked(client, db_session):
    """/rag/query/stream 同样受成本熔断约束，且必须在建立流之前拒绝。"""
    from app.core.config import settings
    from app.models.llm_call_log import LLMCallLog

    old_budget = settings.LLM_BUDGET_CNY
    settings.LLM_BUDGET_CNY = 0.0001
    try:
        client.post(
            "/api/v1/users/",
            json={
                "username": "ragstreambudget",
                "email": "ragstreambudget@test.dev",
                "password": "testpassword",
            },
        )
        login = client.post(
            "/api/v1/auth/login",
            data={"username": "ragstreambudget@test.dev", "password": "testpassword"},
        )
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        user_id = client.get("/api/v1/users/me", headers=headers).json()["id"]

        db_session.add(
            LLMCallLog(
                user_id=user_id,
                endpoint="/api/v1/chat",
                model_name="deepseek-v4.1-flash",
                provider="custom",
                prompt="hi",
                response="x",
                prompt_tokens=1_000_000,
                completion_tokens=0,
                total_tokens=1_000_000,
                latency_ms=1.0,
                estimated_cost_usd=0.0,
                status="success",
            )
        )
        db_session.commit()

        resp = client.post("/api/v1/rag/query/stream", json={"query": "测试"}, headers=headers)
        assert resp.status_code == 402
        assert "成本熔断" in str(resp.json())
    finally:
        settings.LLM_BUDGET_CNY = old_budget


def test_rag_query_stream_cache_hit_skips_llm(client, monkeypatch):
    """流式接口命中语义缓存时直接回放，不产生 LLM 调用。"""
    import app.api.routers.rag as rag_router

    token = _create_user_and_login(client, "stream-cache")

    class FakeEntry:
        answer = "这是缓存里的历史回答"
        original_cost_cny = 0.006

    async def fake_lookup(db, user_id, question):
        return FakeEntry(), 0.98

    called = {"llm": False}

    def fake_get_llm_client(user=None):
        called["llm"] = True
        raise AssertionError("命中缓存时不应获取 LLM 客户端")

    monkeypatch.setattr(rag_router, "lookup_semantic_cache", fake_lookup)
    monkeypatch.setattr(rag_router, "get_llm_client", fake_get_llm_client)

    with client.stream(
        "POST",
        "/api/v1/rag/query/stream",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": "缓存问题", "top_k": 3},
    ) as resp:
        assert resp.status_code == 200
        content = resp.read().decode("utf-8")

    assert "缓存命中" in content
    assert "这是缓存里的历史回答" in content
    assert called["llm"] is False
