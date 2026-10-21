import io

import pytest

from app.utils.file_parser import (
    get_supported_extensions,
    parse_docx,
    parse_file,
    parse_md,
    parse_pdf,
    parse_pptx,
    parse_txt,
)


def _make_docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("第一段：雅思写作任务一")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "词汇"
    table.cell(0, 1).text = "替换表达"
    table.cell(1, 0).text = "important"
    table.cell(1, 1).text = "crucial"
    document.add_paragraph("第三段：结论")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "备考计划"
    slide.placeholders[1].text = "每天背单词 50 个"
    slide.notes_slide.notes_text_frame.text = "记得复习昨天的词"
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(0.8))
    table_shape.table.cell(0, 0).text = "项目"
    table_shape.table.cell(0, 1).text = "目标"
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


class TestParseTxt:
    def test_parse_txt_basic(self):
        result = parse_txt("notes.txt", b"hello world")
        assert result.filename == "notes.txt"
        assert result.file_type == "txt"
        assert result.plain_text == "hello world"

    def test_parse_txt_chinese(self):
        result = parse_txt("doc.txt", "你好，世界！".encode("utf-8"))
        assert result.plain_text == "你好，世界！"


class TestParseMd:
    def test_parse_md_basic(self):
        result = parse_md("readme.md", b"# Title\n\nSome text.")
        assert result.filename == "readme.md"
        assert result.file_type == "md"
        assert "# Title" in result.plain_text

    def test_parse_md_markdown_extension(self):
        result = parse_md("notes.markdown", b"content")
        assert result.file_type == "md"


class TestParsePdf:
    def test_parse_pdf_extracts_text(self):
        pytest.importorskip("fpdf", reason="fpdf2 为测试依赖（requirements-dev.txt），未安装则跳过")
        from fpdf import FPDF

        buf = io.BytesIO()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, text="Project Orion Secret", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(200, 10, text="Code: Aegis-2026-Omega", new_x="LMARGIN", new_y="NEXT")
        pdf.output(buf)
        raw_bytes = buf.getvalue()

        result = parse_pdf("secret.pdf", raw_bytes)
        assert result.file_type == "pdf"
        assert "Project Orion Secret" in result.plain_text
        assert "Aegis-2026-Omega" in result.plain_text


class TestParseDocx:
    def test_extracts_paragraphs_and_tables_in_order(self):
        result = parse_docx("plan.docx", _make_docx_bytes())

        assert result.file_type == "docx"
        assert "第一段：雅思写作任务一" in result.plain_text
        assert "important" in result.plain_text and "crucial" in result.plain_text
        # 表格文字必须夹在两段之间：python-docx 的 paragraphs 不含表格，
        # 若实现成"先所有段落再所有表格"，顺序会错，这里专门钉住
        assert (
            result.plain_text.index("第一段")
            < result.plain_text.index("important")
            < result.plain_text.index("第三段")
        )

    def test_table_row_is_joined(self):
        result = parse_docx("plan.docx", _make_docx_bytes())
        assert "important | crucial" in result.plain_text


class TestParsePptx:
    def test_extracts_slides_tables_and_notes(self):
        result = parse_pptx("plan.pptx", _make_pptx_bytes())

        assert result.file_type == "pptx"
        assert "## 第 1 页" in result.plain_text
        assert "备考计划" in result.plain_text
        assert "每天背单词 50 个" in result.plain_text
        assert "项目 | 目标" in result.plain_text
        assert "[备注] 记得复习昨天的词" in result.plain_text

    def test_empty_presentation_yields_empty_text(self):
        from pptx import Presentation

        buf = io.BytesIO()
        Presentation().save(buf)
        result = parse_pptx("empty.pptx", buf.getvalue())
        assert result.plain_text == ""


class TestParseFileDispatch:
    def test_dispatch_txt(self):
        result = parse_file("doc.txt", b"hi")
        assert result.file_type == "txt"

    def test_dispatch_md(self):
        result = parse_file("doc.md", b"# hi")
        assert result.file_type == "md"

    def test_dispatch_markdown(self):
        result = parse_file("doc.markdown", b"hi")
        assert result.file_type == "md"

    def test_dispatch_pdf(self):
        pytest.importorskip("fpdf", reason="fpdf2 为测试依赖（requirements-dev.txt），未安装则跳过")
        from fpdf import FPDF

        buf = io.BytesIO()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, text="test", new_x="LMARGIN", new_y="NEXT")
        pdf.output(buf)

        result = parse_file("doc.pdf", buf.getvalue())
        assert result.file_type == "pdf"
        assert "test" in result.plain_text

    def test_dispatch_docx(self):
        result = parse_file("plan.docx", _make_docx_bytes())
        assert result.file_type == "docx"

    def test_dispatch_pptx(self):
        result = parse_file("plan.pptx", _make_pptx_bytes())
        assert result.file_type == "pptx"

    def test_dispatch_legacy_doc_gives_actionable_hint(self):
        with pytest.raises(ValueError, match="另存为 .docx"):
            parse_file("old.doc", b"data")

    def test_dispatch_legacy_ppt_gives_actionable_hint(self):
        with pytest.raises(ValueError, match="另存为 .pptx"):
            parse_file("old.ppt", b"data")

    def test_dispatch_unsupported_raises(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            parse_file("image.png", b"data")

    def test_dispatch_case_insensitive(self):
        result = parse_file("DOC.TXT", b"hi")
        assert result.file_type == "txt"


class TestGetSupportedExtensions:
    def test_includes_txt(self):
        assert ".txt" in get_supported_extensions()

    def test_includes_md(self):
        assert ".md" in get_supported_extensions()

    def test_includes_pdf(self):
        assert ".pdf" in get_supported_extensions()

    def test_includes_office_formats(self):
        extensions = get_supported_extensions()
        assert ".docx" in extensions
        assert ".pptx" in extensions
        # 老格式不在支持列表里，否则上传接口会把它放进解析器才报错
        assert ".doc" not in extensions
        assert ".ppt" not in extensions
